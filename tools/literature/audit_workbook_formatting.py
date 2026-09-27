"""Audit the formatting of the literature workbook.

    python tools/literature/audit_workbook_formatting.py

A workbook that is correct but unreadable is not finished. This checks the things
that decide whether a reviewer can use the file: frozen headers, an autofilter, a
sensible column width for every column, a live link wherever a DOI is present, and
no placeholder text left behind. Problems are reported with the sheet and the cell,
and the script exits non-zero when it finds any, so it can gate a commit.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "literature" / "literature_review.xlsx"

# A column narrower than this is almost certainly showing "###" for a number.
MIN_WIDTH = 10
MAX_WIDTH = 60
PLACEHOLDERS = {"tbd", "todo", "n/a", "none", "-", "", "xxx", "???"}


def main() -> int:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl is not installed; cannot audit the workbook", file=sys.stderr)
        return 1
    if not WORKBOOK.exists():
        print(f"missing {WORKBOOK}", file=sys.stderr)
        return 1

    workbook = load_workbook(WORKBOOK)
    problems: list[str] = []
    notes: list[str] = []

    for sheet in workbook.worksheets:
        title = sheet.title
        if sheet.max_row < 2 or sheet.max_column < 1:
            problems.append(f"{title}: sheet is empty")
            continue

        header_row = 1
        # "A2" freezes everything above row two, which is the header row. Any
        # other value means the header scrolls away.
        if sheet.freeze_panes != "A2":
            problems.append(f"{title}: the header row is not frozen "
                            f"(freeze_panes is {sheet.freeze_panes!r})")
        else:
            notes.append(f"{title}: header frozen")

        if sheet.auto_filter.ref is None:
            problems.append(f"{title}: no autofilter on the header row")
        else:
            notes.append(f"{title}: autofilter {sheet.auto_filter.ref}")

        for index in range(1, sheet.max_column + 1):
            letter = sheet.cell(row=header_row, column=index).column_letter
            header = sheet.cell(row=header_row, column=index).value
            if header is None or str(header).strip() == "":
                problems.append(f"{title}!{letter}1: empty column header")

        # Widths are per column, keyed by letter, and openpyxl only stores a width
        # once a column has been touched. An untouched column gets a default, so
        # a missing entry is itself the finding.
        for index in range(1, sheet.max_column + 1):
            letter = sheet.cell(row=header_row, column=index).column_letter
            dimension = sheet.column_dimensions.get(letter)
            width = dimension.width if dimension is not None else None
            header = sheet.cell(row=header_row, column=index).value
            if width is None:
                problems.append(f"{title}!{letter}: column width not set "
                                f"(header {header!r})")
            elif width < MIN_WIDTH:
                problems.append(f"{title}!{letter}: width {width:.1f} is below "
                                f"{MIN_WIDTH} (header {header!r})")
            elif width > MAX_WIDTH:
                notes.append(f"{title}!{letter}: width {width:.1f} exceeds "
                             f"{MAX_WIDTH} (header {header!r})")

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value).strip().lower()
                if text in PLACEHOLDERS and cell.column > 1:
                    problems.append(f"{title}!{cell.coordinate}: placeholder text "
                                    f"{cell.value!r}")

    # Every DOI cell must carry a real hyperlink, not a bare string. The DOI
    # column is looked for on each sheet rather than assumed to be on one, since
    # the workbook keeps the identifier in more than one place.
    doi_columns = 0
    for title in workbook.sheetnames:
        sheet = workbook[title]
        headers = {str(c.value).strip().lower(): c.column
                   for c in sheet[1] if c.value}
        doi_column = next((c for h, c in headers.items()
                           if h in {"doi", "paper link", "paper link (doi)"}), None)
        if doi_column is None:
            continue
        doi_columns += 1
        linked = 0
        present = 0
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=doi_column)
            if not cell.value or not str(cell.value).strip():
                continue
            present += 1
            if cell.hyperlink is not None:
                linked += 1
            else:
                problems.append(
                    f"{title}!{cell.coordinate}: identifier is text without a "
                    f"hyperlink")
        notes.append(f"{title}: {linked}/{present} identifiers hyperlinked")
    if doi_columns == 0:
        problems.append("no sheet carries a DOI or paper-link column")

    for note in notes:
        print(f"note    {note}")
    if problems:
        print()
        for problem in problems:
            print(f"PROBLEM {problem}")
        print(f"\n{len(problems)} formatting problem(s)")
        return 1
    print("\nno formatting problems")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
